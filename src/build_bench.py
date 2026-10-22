# -*- coding: utf-8 -*-
"""經銷商對標資料：讀「2026年經銷商與裕信各廠別損益表」→ 併入 data.json 的 bench
用法：python build_bench.py "對標檔.xlsx"   （需先跑 build.py 產生 data.json）"""
import openpyxl, json, os, sys, shutil, tempfile
HERE = os.path.dirname(os.path.abspath(__file__))
SRC = sys.argv[1] if len(sys.argv) > 1 else r'C:\Users\ES\OneDrive\桌面\AI參照\A谷戰報\2026年裕信各廠別損益表_115.10.06 的複本.xlsx'
tmp = os.path.join(tempfile.gettempdir(), '_bench_src.xlsx'); shutil.copyfile(SRC, tmp)
wb = openpyxl.load_workbook(tmp, data_only=True, read_only=True)
SHEETS = wb.sheetnames
num = lambda x: float(x) if isinstance(x, (int, float)) else None

# ---- 1. 歸類基準（新版檔案沒有此表時，沿用前一版 data.json 內的基準）----
rules = []
if '財務基礎資料歸類基準' in SHEETS:
    for r in wb['財務基礎資料歸類基準'].iter_rows(min_row=3, values_only=True):
        if r[1] or r[2]:
            rules.append(dict(code=r[1] or '', item=(r[2] or '').strip(), formula=r[4] or r[3] or '', define=r[5] or '', review=r[6] or ''))

# ---- 2. 經銷商比例（淨利(占比)簡易解析）----
LINES = [('rev', '零服營收'), ('cost', '零服成本'), ('gp', '零服毛利'), ('bonus', '零服獎金補助'), ('dc', '零服直接成本'), ('ext', '零服外促費用'),
         ('int', '零服內促費用'), ('fixd', '服務廠固定薪(直接)'), ('fixi', '服務廠固定薪(間接)'), ('inv', '零件存貨跌價損失'), ('car', '租金支出-代步車'),
         ('C', '小計'), ('rent', '租金支出-一般.'), ('dep', '折舊與攤提費用'), ('adm', '事務費用'), ('promo', '推廣費用'), ('oth', '其他費用'), ('E', '小計2'),
         ('nm', '零服淨利率'), ('net', '零服淨利 (KNTD)'), ('cars', '進廠台數 (台)-含補償'), ('paid', '進廠台數 (台)-不含補償'), ('perCar', '單台有費營收 (KNTD)')]
ws = wb['淨利(占比)簡易解析']
rows = list(ws.iter_rows(values_only=True))
hdr = rows[1]
DCOL = {}
for ci, h in enumerate(hdr):
    if isinstance(h, str) and ci >= 3:
        nm = h.replace('\n', '').strip()
        if nm in ('全國', '全國(不含裕信)', '裕信', '裕信(台北)', '裕信(新北)', '元隆', '裕新', '裕民', '裕唐', '匯聯', '國通', '裕昌'): DCOL[nm] = ci
NOTE_COL = next(ci for ci, h in enumerate(hdr) if h == '差異說明')
body = [r for r in rows[3:26]]
seen_sub = 0; lab2key = {}
dealers = {k: {} for k in DCOL}
notes = {}
for r in body:
    lab = (r[2] or r[1] or '').strip() if (r[2] or r[1]) else ''
    if lab == '小計':
        seen_sub += 1; lab = '小計' if seen_sub == 1 else '小計2'
    key = next((k for k, l in LINES if l == lab), None)
    if not key: continue
    for d, ci in DCOL.items():
        v = num(r[ci])
        if v is not None: dealers[d][key] = v
    if r[NOTE_COL]: notes[key] = str(r[NOTE_COL])
# 未遮名的差異說明（裕日版）
try:
    ws2 = wb['零服損益比較(裕日)']
    for r in ws2.iter_rows(min_row=9, values_only=True):
        lab = (r[1] or '').strip() if r[1] else ''
        key = next((k for k, l in LINES if l == lab), None)
        if key and r[13]: notes.setdefault(key + '_named', str(r[13]))
except KeyError: pass
# 正負號統一：零服成本為負（對淨利的影響）
for d in dealers.values():
    if 'cost' in d: d['cost'] = -abs(d['cost'])

# ---- 3. 裕信各廠金額（淨利(含金額)）----
ws = wb['淨利(含金額)']
rows = list(ws.iter_rows(values_only=True))
hdr = rows[4]
PLANT_MAP = {'集賢廠': '三重', '撫遠廠': '撫遠鈑噴', '新莊鈑噴廠': '新莊鈑噴', '零服本部(含零件)': '服務部本部', 'DLR 名稱': '裕信合計'}
units = []
for ci in range(3, len(hdr)):
    h = hdr[ci]
    if h:
        nm = PLANT_MAP.get(h, h.replace('廠', '') if h.endswith('廠') else h)
        units.append((ci, nm, h))
PL = {'零服營收': 'rev', '工資營收': 'lab', '零件營收': 'part', '零件營收-外販': 'ext_sale', '零服成本': 'cost', '零服毛利': 'gp', '零服獎金': 'bonus',
      '零服外促費用': 'ext', '零服內促費用': 'int', '服務廠固定薪(直接)': 'fixd', '服務廠固定薪(間接)': 'fixi', '零件存貨跌價損失': 'inv', '租金支出-代步車': 'car',
      '零服直接成本': 'dc', '零服貢獻': 'C', '租金支出-一般.': 'rent', '折舊與攤提費用': 'dep', '事務費用': 'adm', '推廣費用': 'promo', '其他費用': 'oth',
      '零服分攤營運費用': 'E', '零服淨利': 'net', '服務廠進廠數': 'cars', '有費台數(不含補償)': 'paid', '補償台數': 'comp', '服務廠人員數': 'heads',
      '服務廠直接人數': 'dheads', '服務廠間接人數': 'iheads', '直接人員固薪': 'dpay', '間接人員固薪': 'ipay'}
plants = {nm: dict(src=h) for _, nm, h in units}
for r in rows[6:]:
    lab = (r[1] or r[2] or '')
    lab = lab.strip() if isinstance(lab, str) else ''
    key = PL.get(lab)
    if not key: continue
    for ci, nm, _ in units:
        v = num(r[ci])
        if v is not None: plants[nm][key] = v
checks = []
# 合計欄 vs 各單位加總
tot = plants['裕信合計']; parts = [nm for _, nm, _ in units if nm != '裕信合計']
for k in ['rev', 'cost', 'gp', 'bonus', 'dc', 'C', 'rent', 'dep', 'adm', 'oth', 'E', 'net']:
    s = sum(plants[p].get(k, 0) for p in parts)
    checks.append(dict(item=f'裕信合計欄 vs 各單位加總：{k}', book=tot.get(k), calc=s, diff=(tot.get(k) or 0) - s))
# 公式：淨利 = 貢獻 − 分攤
for nm, p in plants.items():
    if 'C' in p and 'E' in p and 'net' in p:
        d = p['net'] - (p['C'] - p['E'])
        if abs(d) > 2: checks.append(dict(item=f'{nm}：零服淨利 ≠ 零服貢獻 − 分攤營運費用', book=p['net'], calc=p['C'] - p['E'], diff=d))
# 更正：合計淨利以「貢獻 − 分攤」為準
tot['net_book'] = tot['net']; tot['net'] = tot['C'] - tot['E']
# 裕信欄（經銷商表）以金額重算，口徑一致
yx = dealers['裕信']; R = tot['rev']
yx_book = dict(yx)
for k in ['cost', 'gp', 'bonus', 'dc', 'ext', 'int', 'fixd', 'fixi', 'inv', 'car', 'C', 'rent', 'dep', 'adm', 'promo', 'oth', 'E']:
    sg = 1 if k in ('gp', 'bonus', 'C') else -1
    yx[k] = sg * tot[k] / R
yx['nm'] = tot['net'] / R; yx['net'] = tot['net']; yx['rev'] = R
if '進廠台數 (台)-含補償':
    checks.append(dict(item='經銷商表「裕信」進廠台數（含補償）', book=yx_book.get('cars'), calc=tot.get('cars'), diff=(yx_book.get('cars') or 0) - (tot.get('cars') or 0)))
    checks.append(dict(item='經銷商表「裕信」進廠台數（不含補償）', book=yx_book.get('paid'), calc=tot.get('paid'), diff=(yx_book.get('paid') or 0) - (tot.get('paid') or 0)))
yx['cars'] = tot.get('cars'); yx['paid'] = tot.get('paid')
checks.append(dict(item='經銷商表「裕信」零服淨利率（原表 vs 重算）', book=yx_book.get('nm'), calc=yx['nm'], diff=(yx_book.get('nm') or 0) - yx['nm']))

# ---- 4. 與儀表板損益資料（2026/1–8）逐廠核對 ----
D = json.load(open(os.path.join(HERE, 'data.json'), encoding='utf-8'))
if not rules and D.get('bench', {}).get('rules'): rules = D['bench']['rules']
ENT = D['meta']['ent']; by = {}
for a in D['acc']: by.setdefault(a['c'], a)
def ours(code, ent, months=range(8)):
    a = by.get(code)
    if not a: return 0
    i = ENT.index(ent); return sum(a['v'][i][m] for m in months) / 1000
for nm, p in plants.items():
    if nm == '裕信合計' or nm not in ENT: continue
    for k, code in [('rev', '100'), ('gp', '120'), ('ext', '140'), ('int', '150'), ('fixd', '16010'), ('fixi', '16020')]:
        o = ours(code, nm)
        if abs(o - p.get(k, 0)) > 1.5: checks.append(dict(item=f'{nm}：{k} 對標檔 vs 儀表板', book=p.get(k), calc=o, diff=(p.get(k) or 0) - o))
checks.append(dict(item='各廠營收、毛利、外促、內促、固薪（直接／間接）與儀表板 2026/1–8 逐項核對', book=None, calc=None, diff=0))

sys.path.insert(0, HERE)
from bench_detail import extract_detail
detail = extract_detail(wb, D, checks)

# ---- 5. 服務部損益說明 V2（正式說明版）：覆蓋 全國／裕信／台北／新北 比例與裕信金額 ----
V2SRC = sys.argv[2] if len(sys.argv) > 2 else r'C:\Users\ES\OneDrive\桌面\AI參照\A谷戰報\服務部損益說明V2 的複本.xlsx'
v2 = None
if os.path.exists(V2SRC):
    from bench_detail import extract_v2
    v2 = extract_v2(V2SRC)
    zs = [r for r in v2['rows'] if r['sec'] == '零服' and r['key']]
    for col, dn in {'nat': '全國', 'natx': '全國(不含裕信)', 'yx': '裕信', 'tp': '裕信(台北)', 'nb': '裕信(新北)',
                    'yx25': '裕信(2025同期)', 'tp25': '裕信(台北)(2025同期)', 'nb25': '裕信(新北)(2025同期)'}.items():
        d = dealers.setdefault(dn, {})
        for r in zs:
            k = r['key']; kv = r['k'].get(col); rv = r['r'].get(col)
            if k == 'rev': d['rev'] = kv if kv is not None else rv
            elif k == 'nm': d['nm'] = rv; d['net'] = kv
            elif rv is not None: d[k] = rv
    for r in v2['rows']:
        if r['n'].startswith('進廠台數') and '含補償' in r['n'] and '不含' not in r['n']:
            for col, dn in {'yx': '裕信', 'tp': '裕信(台北)', 'nb': '裕信(新北)'}.items(): dealers[dn]['cars'] = r['k'][col]
        if '不含補償' in r['n']:
            for col, dn in {'yx': '裕信', 'tp': '裕信(台北)', 'nb': '裕信(新北)'}.items(): dealers[dn]['paid'] = r['k'][col]
        if r['n'].startswith('單台有費營收'):
            for col, dn in {'nat': '全國', 'natx': '全國(不含裕信)', 'yx': '裕信', 'tp': '裕信(台北)', 'nb': '裕信(新北)'}.items():
                if r['k'][col] is not None: dealers[dn]['perCar'] = r['k'][col]
    # 裕信合計金額改用 V2，本部＝合計 − 各廠
    Kyx = {r['key']: r['k']['yx'] for r in zs}
    others = [n for n in plants if n not in ('裕信合計', '服務部本部')]
    for k in ['rent', 'dep', 'adm', 'promo', 'oth', 'car']:
        if Kyx.get(k) is None: continue
        tot[k] = -Kyx[k]
        plants['服務部本部'][k] = tot[k] - sum(plants[n].get(k, 0) for n in others)
    tot['E'] = -Kyx['E']; plants['服務部本部']['E'] = tot['E'] - sum(plants[n].get('E', 0) for n in others)
    tot['net'] = Kyx['nm']; plants['服務部本部']['net'] = plants['服務部本部']['C'] - plants['服務部本部']['E']
    # 檢核：V2 與全科目表（服務部-據點）
    pl = detail.get('pl'); acct = lambda c: next((r['v'][pl['units'].index('合計')] for r in pl['rows'] if r['c'] == c), None) / 1000 if pl else None
    if pl:
        ind = acct('49020'); rent_book = acct('4901010') - acct('611102')
        checks.append(dict(item=f'【V2 已更正】原各廠表「服務部本部」其他費用含間接人事費用（49020，{ind:,.0f}K），不屬零服淨利定義（租金、折舊、事務、推廣、其他）；V2 排除後，裕信零服淨利 {Kyx["nm"]:,.0f}K（{dealers["裕信"]["nm"]*100:.1f}%）。先前以各廠加總 124,011K 更正的作法撤回。', book=None, calc=None, diff=0))
        checks.append(dict(item=f'【待確認】V2 一般租金 {-Kyx["rent"]:,.0f}K，帳上（4901010 − 代步車租金 611102）為 {rent_book:,.0f}K，少 {rent_book + Kyx["rent"]:,.0f}K，剛好等於代步車淨收入 {Kyx["car"]:,.0f}K；而代步車淨收入已計入零服貢獻，疑似重複扣抵，零服淨利可能高估 {rent_book + Kyx["rent"]:,.0f}K', book=rent_book, calc=-Kyx['rent'], diff=rent_book + Kyx['rent']))
    for c in checks:
        if c['item'].startswith('裕信合計欄 vs 各單位加總：net') or c['item'].startswith('裕信合計：零服淨利 ≠'):
            c['item'] += '（原因：本部欄含間接人事費用；V2 已更正，見下方）'
bench = dict(detail=detail, v2=v2, src=os.path.basename(SRC), period='2026/1–8月', unit='千元(K)', lines=LINES, rules=rules, notes=notes,
             dealers=dealers, dealer_book_yx=yx_book, plants=plants, checks=checks)
D['bench'] = bench
json.dump(D, open(os.path.join(HERE, 'data.json'), 'w', encoding='utf-8'), ensure_ascii=False, separators=(',', ':'))
print('dealers', list(dealers.keys()))
print('plants', list(plants.keys()))
for c in checks:
    if c['diff'] and abs(c['diff']) > 0.5: print('CHECK', c)
print('裕信 淨利率 原表', yx_book.get('nm'), '重算', yx['nm'])
