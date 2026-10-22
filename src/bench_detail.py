# -*- coding: utf-8 -*-
"""「2026年裕信各廠別損益表」的說明工作表 → bench.detail（台數、民權UIO、補償、獎金、外促、人事、費用、機具維修、科目×據點）"""
import re

num = lambda x: float(x) if isinstance(x, (int, float)) else None
txt = lambda x: '' if x is None else str(x).strip()


def rows_of(wb, name):
    if name not in wb.sheetnames: return []
    return [list(r) for r in wb[name].iter_rows(values_only=True)]


def cell(r, i):
    return r[i] if i < len(r) else None


def acct_table(rows):
    """科目說明表：表頭列含「科目代號」，每列＝總/細表旗標、科目代號、管理項目、合計、%、備註、改善作法"""
    hi = next((i for i, r in enumerate(rows) if any(txt(v) == '科目代號' for v in r)), None)
    if hi is None: return []
    h = [txt(v) for v in rows[hi]]
    ci = h.index('科目代號'); ni = h.index('管理項目'); ai = h.index('合計')
    pi = h.index('%') if '%' in h else None
    ti = h.index('總/細表') if '總/細表' in h else None
    noi = h.index('備註說明') if '備註說明' in h else None
    aci = h.index('改善作法') if '改善作法' in h else None
    out = []
    for r in rows[hi + 1:]:
        code = txt(cell(r, ci)); nm = txt(cell(r, ni))
        if not code or not nm: continue
        a = num(cell(r, ai))
        out.append(dict(t=txt(cell(r, ti)) == 'Y' if ti is not None else False, c=code, n=nm, a=a or 0,
                        p=num(cell(r, pi)) if pi is not None else None,
                        note=txt(cell(r, noi)) if noi is not None else '', act=txt(cell(r, aci)) if aci is not None else ''))
    # 第一列若沒有標 Y（事務費用表），以「層級最上層」為總計：不另補
    return out


def extract_detail(wb, D, checks):
    det = {}
    ENT = D['meta']['ent']

    # ---------- 1. 台數及單台（定保／維修／鈑噴 × 裕信／新北／台北，2025 vs 2026 同期）----------
    rs = rows_of(wb, '台數及單台'); vol = []
    if rs:
        def block(r0, kind):
            types = [txt(v) for v in rs[r0]]; regs = [txt(v) for v in rs[r0 + 1]]
            y25 = rs[r0 + 2]; y26 = rs[r0 + 3]; cur = ''
            for i in range(1, len(regs)):
                if txt(types[i]): cur = txt(types[i]).replace('台數', '').replace('單台', '')
                if regs[i] and num(cell(y25, i)) is not None:
                    vol.append(dict(k=kind, t=cur, r=regs[i], y25=num(cell(y25, i)), y26=num(cell(y26, i))))
        block(0, 'cnt'); block(8, 'unit')
    det['vol'] = vol

    # ---------- 2. 台數累計統計（全國各經銷商各廠：有費台、定保台 當月／累計 vs 同期）----------
    rs = rows_of(wb, '台數累計統計'); vp = []
    if rs:
        hdr = [txt(v) for v in rs[0]][2:]
        dl = ''
        for r in rs[1:]:
            if txt(r[0]): dl = txt(r[0])
            if not txt(cell(r, 1)): continue
            vp.append(dict(d=dl, p=txt(r[1]), v=[num(x) for x in r[2:2 + len(hdr)]]))
        det['vplants'] = dict(hdr=hdr, rows=vp)

    # ---------- 3. 民權 UIO 現況（民權停業後，客戶轉介各責任廠的保留追蹤）----------
    rs = rows_of(wb, '民權UIO現況')
    if rs:
        blocks = {}; titles = {'推移目標': 'target', '推移實績': 'actual', '推移達成': 'rate', '推移差異': 'diff'}
        for i, r in enumerate(rs):
            t = txt(cell(r, 2))
            if t in titles:
                months = [txt(v) for v in rs[i + 1][3:7]]
                lst = []
                for r2 in rs[i + 2:]:
                    p = txt(cell(r2, 0))
                    if not p or p == '責任廠': break
                    lst.append(dict(p=p, resp=num(cell(r2, 1)), keep=num(cell(r2, 2)), m=[num(x) for x in r2[3:7]]))
                blocks[titles[t]] = dict(months=months, rows=lst)
        # 右側：保留台數（各月新增）、回廠分區、保有／流失 by 經銷商
        side = {}
        for r in rs:
            if txt(cell(r, 8)) == '保留台數': side['keepM'] = [num(x) for x in r[9:14]]
            if txt(cell(r, 8)) == 'UIO數': side['zoneUIO'] = [num(x) for x in r[9:13]]
            if txt(cell(r, 8)) == '有費台數': side['zonePaid'] = [num(x) for x in r[9:13]]
        lost = []
        for r in rs:
            d = txt(cell(r, 15))
            if d.endswith('汽車') or d == '總計':
                lost.append(dict(d=d, keep=num(cell(r, 16)), lost=num(cell(r, 17)), tot=num(cell(r, 18))))
        side['lost'] = lost
        side['uio'] = next((num(cell(r, 18)) for r in rs if txt(cell(r, 17)) == '保有'), None)
        det['uio'] = dict(blocks=blocks, side=side)

    # ---------- 4. 補償台數／專案台數（各經銷商 2026 vs 2025，1–8 月）、補償金額 ----------
    rs = rows_of(wb, '補償台數'); comp = {}
    if rs:
        yr = None
        for r in rs:
            a = txt(cell(r, 0))
            if a in ('2026年', '2025年'): yr = a[:4]; comp[yr] = []; continue
            if a in ('實績', 'NISSAN', '') or yr is None: continue
            vals = [num(x) for x in r[1:20]]
            if all(v is None for v in vals): continue
            comp[yr].append(dict(d=re.sub(r'^\d', '', a), c=vals[0:8], cs=vals[8], p=vals[9:17], ps=vals[17], tot=vals[18]))
    rs2 = rows_of(wb, '補償金額')
    if len(rs2) > 1: comp['amt'] = [num(x) for x in rs2[1][2:10]]
    det['comp'] = comp

    # ---------- 5. 各項科目說明表 ----------
    acc = {}
    for nm in ['零服成本費用說明', '零服獎金', '外促費用說明', '零服內促說明', '固定薪費用說明', '折舊與攤提費用', '事務費用說明', '其他費用說明']:
        acc[nm] = acct_table(rows_of(wb, nm))
    det['acct'] = acc
    # 零服獎金：本部欄
    rs = rows_of(wb, '零服獎金')
    if rs: det['bonusHQ'] = {txt(cell(r, 2)): num(cell(r, 7)) for r in rs if txt(cell(r, 2)) and num(cell(r, 7)) is not None}
    # 外促：YK 護照抵用券項目（裕日補助／裕信吸收）
    rs = rows_of(wb, '外促費用說明'); yk = []
    for r in rs[1:]:
        it = txt(cell(r, 7))
        if it and it != '總計': yk.append(dict(item=it, who=txt(cell(r, 8))))
    det['yk'] = yk
    # 固定薪：平均固薪（2026 vs 2025）
    rs = rows_of(wb, '固定薪費用說明'); pay = []
    for r in rs:
        k = txt(cell(r, 1))
        if k.endswith('人員固薪'): pay.append(dict(k=k, y26=txt(cell(r, 2)), y25=txt(cell(r, 3)), d=txt(cell(r, 4))))
    det['pay'] = pay

    # ---------- 6. 零服獎金細項（大類、規程×細項、入帳月）----------
    rs = rows_of(wb, '零服獎金細項說明'); bcat = []; brule = []
    for r in rs:
        a = txt(cell(r, 0))
        if not a or a.startswith(('零服獎金', '金額', '一、', '二、', '大類', '檢核')): continue
        if len(r) > 8 and txt(cell(r, 1)) and not isinstance(cell(r, 1), (int, float)):
            brule.append(dict(c=a, r=txt(cell(r, 1)), s=txt(cell(r, 2)), n=num(cell(r, 3)), a=num(cell(r, 4)),
                              pc=num(cell(r, 5)), pt=num(cell(r, 6)), avg=num(cell(r, 7)), lag=txt(cell(r, 8))))
        elif num(cell(r, 1)) is not None and a != '合計':
            bcat.append(dict(c=a, n=num(cell(r, 1)), a=num(cell(r, 2)), pt=num(cell(r, 3)), avg=num(cell(r, 4))))
    rs = rows_of(wb, '零服獎金入帳月說明'); bm_cat = []; bm_rule = []; months = []
    for r in rs:
        a = txt(cell(r, 0))
        if a == '大類' and not months: months = [txt(v) for v in r[1:9]]
        if not a or a in ('大類', '合計') or a.startswith(('零服獎金', '依', '一、', '二、', '檢核')): continue
        if isinstance(cell(r, 1), (int, float)):
            bm_cat.append(dict(c=a, m=[num(x) for x in r[1:9]], t=num(cell(r, 9))))
        elif txt(cell(r, 1)):
            bm_rule.append(dict(c=a, r=txt(cell(r, 1)), m=[num(x) for x in r[2:10]], t=num(cell(r, 10))))
    det['bonus'] = dict(cat=bcat, rule=brule, mcat=bm_cat, mrule=bm_rule, months=months)

    # ---------- 7. 外促費用總表（大類／細項 × 月）、成本中心 × 類別 ----------
    rs = rows_of(wb, '外促費用說明-費用總表'); et = []
    for r in rs:
        a = txt(cell(r, 0))
        if not a or a.startswith(('外促費用總表', '成本中心', '金額單位', '大類')): continue
        et.append(dict(c=a, s=txt(cell(r, 1)), m=[num(x) for x in r[2:10]], t=num(cell(r, 10)), p=num(cell(r, 11))))
    det['extTotal'] = et
    # 成本中心 → 單位（以傳票資料對照）
    X = D.get('exp') or {}
    ccmap = {}
    if X:
        f = X['f']; S = X['s']; ic = f.index('cc'); ie = f.index('e')
        from collections import Counter, defaultdict
        cnt = defaultdict(Counter)
        for r in X['rows']: cnt[S[r[ic]]][r[ie]] += 1
        ccmap = {k: ENT[v.most_common(1)[0][0]] for k, v in cnt.items()}
    rs = rows_of(wb, '外促費用說明-成本中心x類別'); ec = None
    for i, r in enumerate(rs):
        if txt(cell(r, 0)) == '成本中心':
            hdr = [txt(v) for v in r[1:] if txt(v)]
            lst = []
            for r2 in rs[i + 1:]:
                cc = txt(cell(r2, 0))
                if not cc: continue
                lst.append(dict(cc=cc, u=('合計' if cc == '總計' else ccmap.get(cc, '')), v=[num(x) for x in r2[1:1 + len(hdr)]]))
            ec = dict(hdr=hdr, rows=lst)
            break
    det['extCC'] = ec

    # ---------- 8. 機具設備維修細項 ----------
    rs = rows_of(wb, '其他費用說明-機具設備維修細項說明'); mm = []; mc = []; mi = []
    sec = 0
    for r in rs:
        a = txt(cell(r, 0))
        if a.startswith('一、'): sec = 1; continue
        if a.startswith('二、'): sec = 2; continue
        if not a or a in ('月份', '大類', '合計', '總計') or a.startswith(('機具維修費用彙整', '資料來源')): continue
        if sec == 1 and a != '合計':
            mm.append(dict(m=num(cell(r, 0)), a=num(cell(r, 1)), n=num(cell(r, 2))))
        elif sec == 2:
            if isinstance(cell(r, 1), (int, float)) and a not in ('合計',):
                mc.append(dict(c=a, m=[num(x) for x in r[1:9]], t=num(cell(r, 9)), p=num(cell(r, 10)), avg=num(cell(r, 11))))
            elif txt(cell(r, 1)) and txt(cell(r, 1)) != '小計' and a != '總計':
                mi.append(dict(c=a, s=txt(cell(r, 1)), m=[num(x) for x in r[2:10]], t=num(cell(r, 10)), n=num(cell(r, 11))))
    det['mach'] = dict(month=mm, cat=mc, item=mi)

    # ---------- 9. 服務部-據點：科目 × 據點（1–8 月，元）----------
    rs = rows_of(wb, '服務部-據點')
    if rs:
        RN = {'撫遠': '撫遠鈑噴', '鈑噴': '新莊鈑噴', '單位：元': '合計'}
        names = [txt(v) for v in rs[1]]
        cols = []
        for i in range(6, len(names)):
            n = names[i] or txt(cell(rs[2], i))
            if not n or n in ('%', '總/細表'): continue
            cols.append((i, RN.get(n, n)))
        body = []
        for r in rs[3:]:
            code = txt(cell(r, 4)); nm = txt(cell(r, 5))
            if not code or not nm: continue
            v = [num(cell(r, i)) or 0 for i, _ in cols]
            body.append(dict(t=txt(cell(r, 3)) == 'Y', c=code, n=nm, v=v))
        keep = [j for j in range(len(cols)) if any(abs(b['v'][j]) > 0.5 for b in body)]
        seen = {}; units = []
        for j in keep:
            n = cols[j][1]; seen[n] = seen.get(n, 0) + 1
            units.append(n if seen[n] == 1 else n + '(' + str(seen[n]) + ')')
        units = [('三重' if u == '三重(2)' else u) for u in units]
        rows = [dict(t=b['t'], c=b['c'], n=b['n'], v=[round(b['v'][j]) for j in keep]) for b in body if any(abs(b['v'][j]) > 0.5 for j in keep)]
        det['pl'] = dict(units=units, rows=rows)
        # 與儀表板 2026/1–8 月逐科目×單位核對（科目代號重複時，任一同代號列吻合即算吻合；零服本部＝服務部本部）
        from collections import defaultdict
        occ = defaultdict(list)
        for a in D['acc']: occ[a['c']].append(a)
        U2 = ['服務部本部' if u == '零服本部' else u for u in units]
        ok = bad = 0; diffs = []
        for r in rows:
            if r['c'] not in occ: continue
            for j, u in enumerate(U2):
                if u not in ENT: continue
                e = ENT.index(u); vals = [sum(a['v'][e][m] for m in range(8)) for a in occ[r['c']]]
                if any(abs(v - r['v'][j]) <= 1 for v in vals): ok += 1
                else:
                    bad += 1; diffs.append(dict(c=r['c'], n=r['n'], u=u, file=r['v'][j], ours=round(vals[0])))
        byu = defaultdict(float)
        for x in diffs:
            if x['c'] == '500': byu[x['u']] += x['file'] - x['ours']
        det['plCheck'] = dict(ok=ok, bad=bad, diffs=diffs, units=sorted({x['u'] for x in diffs}), op=dict(byu))
        net = round(sum(byu.values()))
        extra = [u for u in units if u not in ENT and u not in ('合計', '零服本部') and u not in ('零件', '經理室', '服管課', '行管課', '日技課', '行企課')]
        op_row = next((r for r in rows if r['c'] == '500'), None)
        ex_op = {u: op_row['v'][units.index(u)] for u in extra} if op_row else {}
        det['plCheck']['extraUnits'] = ex_op
        checks.append(dict(item=f'服務部-據點（科目×單位，1–8 月）vs 儀表板損益表：吻合 {ok}／{ok + bad}；不符僅在 {"、".join(sorted({x["u"] for x in diffs}))}。'
                                f'原因：新檔把撫遠鈑噴 1 月記在本部的費用改歸撫遠（撫遠營業利益 {round(byu.get("撫遠鈑噴", 0))} 元），'
                                f'並另列單位 {"、".join(f"{u}（營業利益 {v:,.0f} 元）" for u, v in ex_op.items())}；兩者合計等於本部的差異 {round(byu.get("服務部本部", 0)):,} 元，服務部合計不變',
                           book=ok + bad, calc=ok, diff=bad))
    return det


V2_KEYS = {'零服營收': 'rev', '零服成本': 'cost', '零服毛利': 'gp', '零服獎金補助': 'bonus', '零服外促費用': 'ext', '零服內促費用': 'int',
           '服務廠固定薪(直接)': 'fixd', '服務廠固定薪(間接)': 'fixi', '零件存貨跌價損失': 'inv', '租金支出-代步車': 'car', '租金支出-一般.': 'rent',
           '折舊與攤提費用': 'dep', '事務費用': 'adm', '推廣費用': 'promo', '零服淨利': 'nm', '零服吸收率(%)': 'absorb', '零服直接成本': 'dc'}


def extract_v2(path):
    """服務部損益說明 V2：逐科目 全國／全國(不含裕信)／裕信／台北／新北／最優 DLR／去年同期 + 專責部門、優劣、說明、對策、內部備註"""
    import openpyxl, shutil, tempfile, os
    tmp = os.path.join(tempfile.gettempdir(), '_v2_src.xlsx'); shutil.copyfile(path, tmp)
    wb = openpyxl.load_workbook(tmp, data_only=True, read_only=True)
    rs = [list(r) for r in wb[wb.sheetnames[0]].iter_rows(values_only=True)]
    title = txt(rs[0][1]).split('\n')[0] if rs else ''
    out = []; sec = '新車與周邊'; grp = ''; sub_n = 0
    for r in rs[2:]:
        c1, c2 = txt(cell(r, 1)), txt(cell(r, 2))
        lab = c2 or c1
        if not lab or lab.endswith('!A1'):
            lab = c2 if c2 and not c2.endswith('!A1') else (txt(cell(r, 3)) if c2.endswith('!A1') else lab)
        if c2.endswith('!A1'): lab = txt(cell(r, 3))
        if not lab: continue
        if lab == '保有UIO數(台)': sec = '服務量能'
        if lab == '零服營收': sec = '零服'; sub_n = 0
        if c2 and c1: grp = c1
        elif c1 and not c2: grp = ''
        key = V2_KEYS.get(lab)
        if sec == '零服' and lab == '小計':
            sub_n += 1; key = 'C' if sub_n == 1 else 'E'
        if sec == '零服' and lab.startswith('其他費用'): key = 'oth'
        n = lambda i: num(cell(r, i))
        if sec != '零服': key = None
        if grp.endswith('!A1'): grp = ''
        if all(n(i) is None for i in (3, 4, 5, 6, 8, 9, 15, 17, 35)): continue
        ex = txt(cell(r, 12)) == '優/劣'  # 原表第一列為填寫範例
        out.append(dict(sec=sec, grp=grp, n=lab, key=key, dir=txt(cell(r, 0)),
                        k=dict(nat=n(3), natx=n(5), yx=n(8), tp=n(15), nb=n(17), yx25=n(35), tp25=n(37), nb25=n(39)),
                        r=dict(nat=n(4), natx=n(6), yx=n(9), tp=n(16), nb=n(18), best=n(19), yx25=n(36), tp25=n(38), nb25=n(40)),
                        best=txt(cell(r, 20)),
                        d=dict(yx=[n(23), n(24), n(25)], tp=[n(27), n(28), n(29)], nb=[n(31), n(32), n(33)]),
                        dept=txt(cell(r, 10)), factor='' if ex else txt(cell(r, 11)), judge='' if ex else txt(cell(r, 12)), expl='' if ex else txt(cell(r, 13)), act='' if ex else txt(cell(r, 14)),
                        internal=txt(cell(r, 41)), extra=txt(cell(r, 42))))
    return dict(src=os.path.basename(path), title=title, rows=out)
